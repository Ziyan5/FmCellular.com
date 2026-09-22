#!/usr/bin/env python3
"""Read-only exporter for FM Cellular's approved public workbook sheets."""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


APPROVED_SHEETS = {
    "Inventory": "inventory.csv",
    "Parts Catalogue": "parts_catalogue.csv",
    "Photo Library": "photo_library.csv",
    "Finishes": "finishes.csv",
    "Site text": "site_text.csv",
    "Posters": "posters.csv",
    "Buyback": "buyback_config.csv",
}


def clean_header(value: object, index: int) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text or f"column_{index + 1}"


def export_sheet(workbook, sheet_name: str, output_path: Path) -> int:
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        return 0

    headers = [clean_header(value, i) for i, value in enumerate(header_row)]
    last_named = max((i for i, name in enumerate(headers) if not name.startswith("column_")), default=-1)
    headers = headers[: last_named + 1]

    written = 0
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            values = list(row[: len(headers)])
            if not any(value is not None and str(value).strip() for value in values):
                continue
            writer.writerow(["" if value is None else value for value in values])
            written += 1
    return written


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: export_public_catalog.py INPUT.xlsx OUTPUT_DIRECTORY", file=sys.stderr)
        return 2

    source = Path(sys.argv[1]).expanduser().resolve()
    destination = Path(sys.argv[2]).expanduser().resolve()
    if not source.is_file():
        print(f"Workbook not found: {source}", file=sys.stderr)
        return 2

    destination.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    missing = sorted(set(APPROVED_SHEETS) - set(workbook.sheetnames))
    if missing:
        print(f"Missing approved sheets: {', '.join(missing)}", file=sys.stderr)
        return 1

    total = 0
    for sheet_name, filename in APPROVED_SHEETS.items():
        count = export_sheet(workbook, sheet_name, destination / filename)
        total += count
        print(f"{filename}: {count} data rows")
    print(f"Exported {total} public rows. The source workbook was not modified.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
