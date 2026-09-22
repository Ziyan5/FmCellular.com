#!/usr/bin/env python3
"""Build idempotent Supabase seed SQL from public Master List sheets."""

from __future__ import annotations

import hashlib
import json
import re
import sys
import uuid
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

NAMESPACE = uuid.UUID("fbc0be6e-87d2-4f90-8d94-f2114fa15e9e")
BATCH_SIZE = 2000


def text(value):
    return str(value).strip() if value is not None else ""


def money(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace("$", "").replace(",", "").strip()).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def integer(value):
    try:
        return max(0, int(float(value))) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def yes(value):
    return text(value).lower() in {"1", "true", "yes", "y", "in stock", "available"}


def q(value):
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def slugify(value):
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug[:110] or "item"


def brand_for(category, model):
    combined = f"{category} {model}".lower()
    if any(word in combined for word in ("iphone", "ipad", "macbook", "apple watch", "airpod")):
        return "Apple"
    if any(word in combined for word in ("samsung", "galaxy", "z fold", "z flip")):
        return "Samsung"
    if "google" in combined or "pixel" in combined:
        return "Google"
    return None


def uid(kind, key):
    return str(uuid.uuid5(NAMESPACE, f"{kind}|{key}"))


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    header = [text(v).lower() for v in next(iterator)]
    for row_number, values in enumerate(iterator, 2):
        record = {header[i]: values[i] for i in range(min(len(header), len(values))) if header[i]}
        if any(text(v) for v in record.values()):
            yield row_number, record


def make_statements(workbook):
    products = {}
    variants = []
    duplicate_counter = defaultdict(int)

    for sheet_name, item_type in (("Inventory", "device"), ("Parts Catalogue", "part")):
        for row_number, record in rows(workbook[sheet_name]):
            category, model = text(record.get("category")), text(record.get("model"))
            part_name = text(record.get("part")) or None
            if not category or not model:
                continue
            product_key = "|".join((item_type, category.lower(), model.lower(), (part_name or "").lower()))
            product_id = uid("product", product_key)
            if product_id not in products:
                base_slug = slugify("-".join(filter(None, (category, model, part_name))))
                slug = f"{base_slug}-{hashlib.sha1(product_key.encode()).hexdigest()[:8]}"
                products[product_id] = (product_id, item_type, category, brand_for(category, model), model, part_name, slug)

            storage = text(record.get("storage")) or None
            condition = text(record.get("condition")) or None
            color = text(record.get("color")) or None
            supplier_sku = text(record.get("supplier_sku")) or None
            raw_key = "|".join(map(str, (product_key, storage or "", condition or "", color or "", supplier_sku or "", money(record.get("price")))))
            duplicate_counter[raw_key] += 1
            variant_key = f"{raw_key}|{duplicate_counter[raw_key]}"
            variant_id = uid("variant", variant_key)
            sku = "FMC-" + hashlib.sha1(variant_key.encode()).hexdigest()[:16].upper()
            quantity = integer(record.get("qty"))
            available = quantity > 0 if quantity is not None else yes(record.get("in_stock"))
            variants.append({
                "id": variant_id, "product_id": product_id, "sku": sku, "supplier_sku": supplier_sku,
                "storage": storage, "condition": condition, "color": color,
                "retail_price": money(record.get("price")), "msrp": money(record.get("msrp")),
                "quantity": quantity, "available": available,
                "wholesale": money(record.get("wholesale")), "image_url": text(record.get("image_url")) or None,
                "source": sheet_name, "source_row": row_number,
            })

    statements = []
    for p in products.values():
        statements.append(
            "insert into public.catalog_products (id,item_type,category,brand,model,part_name,slug) values "
            f"({q(p[0])}::uuid,{q(p[1])}::public.catalog_item_type,{q(p[2])},{q(p[3])},{q(p[4])},{q(p[5])},{q(p[6])}) "
            "on conflict (id) do update set category=excluded.category,brand=excluded.brand,model=excluded.model,part_name=excluded.part_name,slug=excluded.slug;"
        )
    for v in variants:
        metadata = json.dumps({"source_sheet": v["source"], "source_row": v["source_row"]}, separators=(",", ":"))
        statements.append(
            "insert into public.catalog_variants (id,product_id,sku,supplier_sku,storage,condition,color,retail_price,msrp,metadata) values "
            f"({q(v['id'])}::uuid,{q(v['product_id'])}::uuid,{q(v['sku'])},{q(v['supplier_sku'])},{q(v['storage'])},{q(v['condition'])},{q(v['color'])},{q(v['retail_price'])}::numeric,{q(v['msrp'])}::numeric,{q(metadata)}::jsonb) "
            "on conflict (id) do update set supplier_sku=excluded.supplier_sku,storage=excluded.storage,condition=excluded.condition,color=excluded.color,retail_price=excluded.retail_price,msrp=excluded.msrp,metadata=excluded.metadata;"
        )
        statements.append(
            "insert into public.inventory_levels (variant_id,quantity,available) values "
            f"({q(v['id'])}::uuid,{q(v['quantity'])}::integer,{str(v['available']).lower()}) "
            "on conflict (variant_id) do update set quantity=excluded.quantity,available=excluded.available;"
        )
        if v["wholesale"] is not None:
            statements.append(
                "insert into public.wholesale_prices (variant_id,price) values "
                f"({q(v['id'])}::uuid,{q(v['wholesale'])}::numeric) "
                "on conflict (variant_id) do update set price=excluded.price;"
            )
        if v["image_url"] and v["image_url"].startswith(("https://", "http://")):
            statements.append(
                "insert into public.product_images (product_id,variant_id,url,is_primary) values "
                f"({q(v['product_id'])}::uuid,{q(v['id'])}::uuid,{q(v['image_url'])},true) "
                "on conflict (product_id,url) do nothing;"
            )
    return products, variants, statements


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_seed_sql.py INPUT.xlsx OUTPUT_DIRECTORY")
    source, output = Path(sys.argv[1]), Path(sys.argv[2])
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    products, variants, statements = make_statements(workbook)
    for index in range(0, len(statements), BATCH_SIZE):
        batch = statements[index:index + BATCH_SIZE]
        path = output / f"seed_{index // BATCH_SIZE + 1:03d}.sql"
        path.write_text("begin;\n" + "\n".join(batch) + "\ncommit;\n", encoding="utf-8")
    print(json.dumps({"products": len(products), "variants": len(variants), "statements": len(statements), "batches": (len(statements)+BATCH_SIZE-1)//BATCH_SIZE}))


if __name__ == "__main__":
    main()
